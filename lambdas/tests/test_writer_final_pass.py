"""UOW-D — the human downloads the workbook they edited.

The writer used to run once, before the review gate, so the workbook the human
downloaded at step 6 predated their own edits at step 5. It now runs twice, and
the answers come from the Questions table on both passes — which is why the SME
edits review_api writes there need no special handling in the writer.
"""

from __future__ import annotations

import tempfile
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import boto3
import openpyxl
import pytest
from moto import mock_aws

from shared import session

FIXTURE = Path(__file__).resolve().parents[2] / "data" / "incoming" / "sample_rfp_acmesec.xlsx"

JOBS = "test-jobs"
QUESTIONS = "test-questions"
REVIEWS = "test-reviews"
INCOMING = "test-incoming"
OUTPUT = "test-output"

MACHINE_ANSWER = "Keys are rotated on a schedule defined by policy."
SME_ANSWER = "Customer-managed keys are rotated every 90 days, verified by SOC 2 CC6.1."


def _question(question_id: str, answer_text: str, *, sme_edited: bool, row: int) -> dict[str, Any]:
    return {
        "jobId": "job-1",
        "questionId": question_id,
        "text": "How often are encryption keys rotated?",
        "answer_text": answer_text,
        "sme_edited": sme_edited,
        "raw_confidence": Decimal("0.82"),
        "tier": "green" if sme_edited else "amber",
        "confidence_breakdown": {
            "h": Decimal("0.8"),
            "r": Decimal("0.8"),
            "c": Decimal("0.8"),
            "f": Decimal("0.8"),
            "g": Decimal("1.0"),
        },
        "hard_rule_triggers": [],
        "reviewer_required": False,
        "answer_cell": {"sheet": "Questionnaire", "coordinate": f"C{row}"},
        "confidence_cell": {"sheet": "Questionnaire", "coordinate": f"D{row}"},
    }


@pytest.fixture
def writer(monkeypatch: pytest.MonkeyPatch) -> Any:
    from excel_writer import handler as writer_handler

    with mock_aws():
        ddb = boto3.resource("dynamodb", region_name="us-east-1")
        s3 = boto3.client("s3", region_name="us-east-1")
        ddb.create_table(
            TableName=JOBS,
            KeySchema=[{"AttributeName": "jobId", "KeyType": "HASH"}],
            AttributeDefinitions=[{"AttributeName": "jobId", "AttributeType": "S"}],
            BillingMode="PAY_PER_REQUEST",
        )
        for name, sort_key in ((QUESTIONS, "questionId"), (REVIEWS, "reviewedAt")):
            ddb.create_table(
                TableName=name,
                KeySchema=[
                    {"AttributeName": "jobId", "KeyType": "HASH"},
                    {"AttributeName": sort_key, "KeyType": "RANGE"},
                ],
                AttributeDefinitions=[
                    {"AttributeName": "jobId", "AttributeType": "S"},
                    {"AttributeName": sort_key, "AttributeType": "S"},
                ],
                BillingMode="PAY_PER_REQUEST",
            )
        for bucket in (INCOMING, OUTPUT):
            s3.create_bucket(Bucket=bucket)
        s3.upload_file(str(FIXTURE), INCOMING, "incoming/job-1/caiq.xlsx")

        for key, value in {
            "JOBS_TABLE": JOBS,
            "QUESTIONS_TABLE": QUESTIONS,
            "REVIEWS_TABLE": REVIEWS,
            "OUTPUT_BUCKET": OUTPUT,
        }.items():
            monkeypatch.setenv(key, value)
        monkeypatch.setattr(writer_handler, "ddb", ddb)
        monkeypatch.setattr(writer_handler, "s3", s3)
        monkeypatch.setattr(session, "_ddb_resource", ddb)

        ddb.Table(QUESTIONS).put_item(Item=_question("q-0001", SME_ANSWER, sme_edited=True, row=8))
        ddb.Table(QUESTIONS).put_item(
            Item=_question("q-0002", MACHINE_ANSWER, sme_edited=False, row=9)
        )
        ddb.Table(JOBS).put_item(
            Item={"jobId": "job-1", "session_status": session.APPROVED, "review_status": "APPROVED"}
        )
        yield SimpleNamespace(handler=writer_handler, ddb=ddb, s3=s3)


def _invoke(writer: Any, stage: str) -> Any:
    return writer.handler.lambda_handler(
        {
            "jobId": "job-1",
            "sourceBucket": INCOMING,
            "sourceKey": "incoming/job-1/caiq.xlsx",
            "stage": stage,
        },
        None,
    )


def _cell_values(writer: Any, output_key: str) -> list[str]:
    with tempfile.TemporaryDirectory() as tmp:
        dst = Path(tmp) / "out.xlsx"
        writer.s3.download_file(OUTPUT, output_key, str(dst))
        ws = openpyxl.load_workbook(str(dst))["Questionnaire"]
        return [str(ws["C8"].value), str(ws["C9"].value)]


def test_final_workbook_carries_the_sme_edit(writer: Any) -> None:
    result = _invoke(writer, "final")

    values = _cell_values(writer, result["outputKey"])
    assert values[0] == SME_ANSWER
    assert values[1] == MACHINE_ANSWER


def test_draft_and_final_are_different_objects(writer: Any) -> None:
    draft = _invoke(writer, "draft")
    final = _invoke(writer, "final")

    assert draft["outputKey"] != final["outputKey"]
    assert draft["outputKey"].endswith("_answered.xlsx")
    assert final["outputKey"].endswith("_final.xlsx")


def test_stage_defaults_to_draft(writer: Any) -> None:
    """An invocation without a stage must not be mistaken for the final pass."""
    result = writer.handler.lambda_handler(
        {"jobId": "job-1", "sourceBucket": INCOMING, "sourceKey": "incoming/job-1/caiq.xlsx"},
        None,
    )
    assert result["stage"] == "draft"
    assert (
        writer.ddb.Table(JOBS).get_item(Key={"jobId": "job-1"})["Item"]["session_status"]
        == session.APPROVED
    )


def test_draft_pass_records_tier_counts_but_not_completion(writer: Any) -> None:
    _invoke(writer, "draft")

    job = writer.ddb.Table(JOBS).get_item(Key={"jobId": "job-1"})["Item"]
    assert {k: int(v) for k, v in job["tier_counts"].items()} == {"green": 1, "amber": 1, "red": 0}
    assert job["session_status"] == session.APPROVED  # still not downloadable


def test_final_pass_completes_the_session_and_moves_the_download(writer: Any) -> None:
    result = _invoke(writer, "final")

    job = writer.ddb.Table(JOBS).get_item(Key={"jobId": "job-1"})["Item"]
    assert job["session_status"] == session.COMPLETE
    # The download route signs job["output_key"]; it now points at the final file.
    assert job["output_key"] == result["outputKey"]
    assert job["output_bucket"] == OUTPUT

    reviews = writer.ddb.Table(REVIEWS).query(
        KeyConditionExpression=boto3.dynamodb.conditions.Key("jobId").eq("job-1")
    )["Items"]
    assert [r["step"] for r in reviews] == [session.COMPLETE]
