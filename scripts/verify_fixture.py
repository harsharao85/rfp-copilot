#!/usr/bin/env python3
"""Offline verification of the Echoworx CAIQ fixture against the real pipeline.

Runs every question in data/incoming/echoworx-caiq-v4-40q.xlsx through the
shipped question_classifier, confidence_scorer and hard_rules code with a
stubbed retriever, and checks the outcome against
data/incoming/echoworx-caiq-v4-40q.expected.json.

No AWS. Every AWS endpoint is firewalled in this environment, so nothing here
may construct a boto3 client, call Bedrock, or read DynamoDB.

Usage:
  PYTHONPATH=. python scripts/verify_fixture.py [--verbose] [--json OUT.json]

What is real and what is stubbed
--------------------------------
Real, imported from lambdas/ and not reimplemented:
  question_classifier.handler._extract_topics   topic keyword matching
  question_classifier.dispatch.get_dispatch_plan  class + source routing
  confidence_scorer.scorer.score                composite and class tier logic
  hard_rules.rules.apply_rules                  the six rules and their triggers
  shared.provenance.Provenance                  sidecar hydration, expiry, citation

Stubbed, because the real ones need AWS:

  Retrieval. bedrock-agent-runtime:Retrieve is replaced by TF-IDF cosine over
  the same corpus text that scripts/generate_corpus_documents.py writes, split
  into passages on document headings, and filtered by source_type exactly as
  retriever._SOURCE_TYPE_FILTER does. Cosine is then mapped onto the score
  range an embedding retriever actually returns - a lexical cosine of 0 is not
  a semantic similarity of 0 - by SCORE_FLOOR + SCORE_SPAN * cosine. The
  mapping is a calibration of a lexical proxy, and it is the one number in this
  harness that is a judgement call rather than a fact; it is stated here rather
  than buried so that a reader can discount it.

  CustomerRefs. The DynamoDB scan is replaced by the same filter over
  data/graph/customers.json that seeds the table: public_reference is true AND
  approval_expires is not before today.

  Generation. Bedrock is replaced by an extractive composer: restate the
  question, then append the top grounding passages. Restating is not a
  shortcut - hard_rules matches on the ANSWER, and a generator that is asked
  "will you support X by Q2 2027" produces an answer carrying that framing. If
  the stub dropped the restatement, questions 31 to 37 would silently stop
  firing the rules they exist to fire, and the fixture would prove nothing.
"""
from __future__ import annotations

import argparse
import json
import math
import re
import sys
from collections import Counter
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
LAMBDAS = ROOT / "lambdas"
# question_classifier goes on the path because its handler does `from dispatch
# import ...` the way it does inside the Lambda bundle. hard_rules deliberately
# does NOT: it also has a module called `handler`, and putting both on the path
# makes `import handler` ambiguous — and the losing import constructs a
# DynamoDB resource at module scope, which is exactly what must not happen here.
for extra in (LAMBDAS, LAMBDAS / "question_classifier"):
    sys.path.insert(0, str(extra))

from dispatch import get_dispatch_plan  # noqa: E402
from handler import _extract_topics  # noqa: E402  question_classifier/handler.py

from confidence_scorer.scorer import score as composite_score  # noqa: E402
from shared.models import GeneratedAnswer, PriorAnswerMatch, RetrievedPassage, Tier  # noqa: E402
from shared.provenance import Provenance, render_citation  # noqa: E402


def _load_rules() -> Any:
    """Load lambdas/hard_rules/rules.py by path, under a name that cannot
    collide with the classifier's modules."""
    import importlib.util  # noqa: PLC0415

    spec = importlib.util.spec_from_file_location(
        "_hard_rules_rules", LAMBDAS / "hard_rules" / "rules.py"
    )
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


hard_rules = _load_rules()

sys.path.insert(0, str(ROOT / "scripts"))
import generate_corpus_documents as corpusgen  # noqa: E402

DATA = ROOT / "data"
CORPUS_REAL = DATA / "corpus-real"
FIXTURE = DATA / "incoming" / "echoworx-caiq-v4-40q.xlsx"
MANIFEST = DATA / "incoming" / "echoworx-caiq-v4-40q.expected.json"

# ---------------------------------------------------------------------------
# Retrieval stub calibration. See the module docstring.
# ---------------------------------------------------------------------------
SCORE_FLOOR = 0.30   # what an embedding retriever returns for unrelated text
SCORE_SPAN = 0.68    # so a perfect lexical match maps to 0.98
MIN_COSINE = 0.045   # below this the passage is not returned at all
TOP_K = 5            # per source, matching retriever._knowledge_base_retrieve

# retriever._SOURCE_TYPE_FILTER / _SOURCE_TO_SYSTEM, mirrored rather than
# imported because importing lambdas/retriever/handler.py pulls in boto3.
SOURCE_TYPE_FILTER = {
    "compliance_store": "compliance_cert",
    "prior_rfps": "prior_rfp",
    "product_docs": "product_doc",
    "sme_approved_answers": "sme_approved_answer",
}
SOURCE_TO_SYSTEM = {
    "compliance_store": "compliance",
    "prior_rfps": "historical_rfp",
    "product_docs": "whitepaper",
    "sme_approved_answers": "sme_approved",
}
PRIMARY_SYSTEMS = {"compliance", "whitepaper"}

_TOKEN_RE = re.compile(r"[a-z0-9]+")
_STOPWORDS = frozenset(
    """a an the and or of to in for on at by is are be been with from as that this
    these those it its do does did you your we our us they their there here what
    which who whom how why when where can could will would shall should may might
    must not no if then than so such into over under about per each any all both
    more most other some only own same too very just also state provide describe
    confirm including include list detail" """.split()
)


def _tokens(text: str) -> list[str]:
    return [t for t in _TOKEN_RE.findall(text.lower()) if t not in _STOPWORDS and len(t) > 2]


# ---------------------------------------------------------------------------
# Corpus index
# ---------------------------------------------------------------------------


@dataclass
class Chunk:
    document_id: str
    source_type: str
    heading: str
    text: str
    sidecar: dict[str, Any]
    uri: str
    tokens: Counter = field(default_factory=Counter)


def _split_markdown(body: str) -> list[tuple[str, str]]:
    """Split on '## ' headings. Returns [(heading, text)]."""
    parts: list[tuple[str, str]] = []
    heading = "Overview"
    buf: list[str] = []
    for line in body.splitlines():
        if line.startswith("## "):
            if buf and any(b.strip() for b in buf):
                parts.append((heading, "\n".join(buf).strip()))
            heading = line[3:].strip()
            buf = []
        elif line.startswith("# "):
            heading = line[2:].strip()
            buf = []
        else:
            buf.append(line)
    if buf and any(b.strip() for b in buf):
        parts.append((heading, "\n".join(buf).strip()))
    return parts


def _split_pdf_blocks(blocks: list[dict[str, Any]]) -> list[tuple[str, str]]:
    """Group reportlab blocks under their preceding heading."""
    parts: list[tuple[str, str]] = []
    heading = "Overview"
    buf: list[str] = []
    for block in blocks:
        if "h" in block:
            if buf:
                parts.append((heading, "\n".join(buf).strip()))
            heading = block["h"]
            buf = []
        elif "p" in block:
            buf.append(block["p"])
        elif "table" in block:
            for row in block["table"]:
                buf.append(" | ".join(str(c) for c in row))
    if buf:
        parts.append((heading, "\n".join(buf).strip()))
    return parts


def _sidecar_for(prefix: str, document_id: str, ext: str) -> dict[str, Any]:
    path = CORPUS_REAL / prefix / f"{document_id}{ext}.metadata.json"
    if not path.exists():
        raise SystemExit(f"missing sidecar for {prefix}/{document_id}: {path}")
    return json.loads(path.read_text())["metadataAttributes"]


def build_index() -> list[Chunk]:
    """Reconstruct passage text from the same bodies the generator renders.

    Reading the compliance PDFs back would need a PDF text extractor and would
    reintroduce layout noise the embedder never sees, so the blocks are read
    from generate_corpus_documents directly. Metadata always comes from the
    sidecar on disk, so a sidecar that was never written is a hard failure here
    rather than a silent gap.
    """
    chunks: list[Chunk] = []

    def add(prefix: str, doc_id: str, ext: str, parts: list[tuple[str, str]]) -> None:
        sidecar = _sidecar_for(prefix, doc_id, ext)
        uri = f"s3://reference-corpus/{prefix}/{doc_id}{ext}"
        for heading, text in parts:
            if not text.strip():
                continue
            chunks.append(
                Chunk(
                    document_id=doc_id,
                    source_type=sidecar["source_type"],
                    heading=heading,
                    text=text,
                    sidecar=sidecar,
                    uri=uri,
                    tokens=Counter(_tokens(f"{heading} {text}")),
                )
            )

    for doc_id, blocks in corpusgen.COMPLIANCE_BODIES.items():
        add("compliance", doc_id, ".pdf", _split_pdf_blocks(blocks))
    for doc_id, body in corpusgen.PRODUCT_DOCS_BODIES.items():
        add("product-docs", doc_id, ".md", _split_markdown(body))
    for doc_id, body in corpusgen.PRIOR_RFP_BODIES.items():
        add("prior-rfps", doc_id, ".md", _split_markdown(body))

    for src in sorted((DATA / "corpus" / "sme-approved").glob("*.json")):
        data = json.loads(src.read_text())
        add(
            "sme-approved",
            data["document_id"],
            ".md",
            [("Question and answer", f"{data['question_text']}\n\n{data['answer_text']}")],
        )

    return chunks


def build_idf(chunks: list[Chunk]) -> dict[str, float]:
    n = len(chunks)
    df: Counter[str] = Counter()
    for c in chunks:
        df.update(c.tokens.keys())
    return {term: math.log((n + 1) / (count + 1)) + 1.0 for term, count in df.items()}


def _cosine(q: Counter[str], d: Counter[str], idf: dict[str, float]) -> float:
    qv = {t: (1 + math.log(c)) * idf.get(t, 1.0) for t, c in q.items()}
    dv = {t: (1 + math.log(c)) * idf.get(t, 1.0) for t, c in d.items()}
    qn = math.sqrt(sum(v * v for v in qv.values()))
    dn = math.sqrt(sum(v * v for v in dv.values()))
    if qn == 0 or dn == 0:
        return 0.0
    dot = sum(v * dv.get(t, 0.0) for t, v in qv.items())
    return dot / (qn * dn)


# ---------------------------------------------------------------------------
# Stub retriever
# ---------------------------------------------------------------------------


class StubRetriever:
    def __init__(self, chunks: list[Chunk]) -> None:
        self.chunks = chunks
        self.idf = build_idf(chunks)

    def retrieve(self, query: str, source: str, top_k: int = TOP_K) -> list[RetrievedPassage]:
        source_type = SOURCE_TYPE_FILTER.get(source)
        if source_type is None:
            return []
        system = SOURCE_TO_SYSTEM[source]
        qtok = Counter(_tokens(query))
        scored: list[tuple[float, Chunk]] = []
        for chunk in self.chunks:
            if chunk.source_type != source_type:
                continue
            cos = _cosine(qtok, chunk.tokens, self.idf)
            if cos >= MIN_COSINE:
                scored.append((cos, chunk))
        scored.sort(key=lambda pair: pair[0], reverse=True)

        passages: list[RetrievedPassage] = []
        for cos, chunk in scored[:top_k]:
            meta: dict[str, Any] = {
                k: v for k, v in chunk.sidecar.items()
                if isinstance(v, (str, int, float, bool))
                or (isinstance(v, list) and all(isinstance(x, str) for x in v))
            }
            meta["s3_uri"] = chunk.uri
            try:
                provenance = Provenance.from_sidecar(dict(chunk.sidecar))
            except Exception:  # noqa: BLE001 — mirrors retriever's fail-soft hydration
                provenance = None
            passages.append(
                RetrievedPassage(
                    source_system=system,  # type: ignore[arg-type]
                    document_id=chunk.document_id,
                    excerpt=f"{chunk.heading}. {chunk.text}"[:2000],
                    score=round(min(SCORE_FLOOR + SCORE_SPAN * cos, 1.0), 4),
                    uri=chunk.uri,
                    metadata=meta,
                    provenance=provenance,
                )
            )
        return passages

    def prior_matches(self, query: str, top_k: int = 3) -> list[PriorAnswerMatch]:
        """Mirror of retriever._kb_prior_matches, including the lexical
        expires_on comparison at the retrieval boundary."""
        today = date.today().isoformat()
        out: list[PriorAnswerMatch] = []
        for p in self.retrieve(query, "sme_approved_answers", top_k=top_k):
            md = p.metadata
            expires_on = str(md.get("expires_on") or "")
            if expires_on and expires_on < today:
                continue
            out.append(
                PriorAnswerMatch(
                    answer_id=p.document_id,
                    question_text=str(md.get("question_text", "")) or p.excerpt[:200],
                    answer_text=p.excerpt,
                    similarity=p.score,
                    approved_by=str(md.get("approved_by", "unknown")),
                    approved_at=datetime.fromisoformat(
                        str(md.get("approved_at", "2025-01-01")).replace("Z", "+00:00")
                    ),
                    expires_on=None,
                )
            )
        return out


# ---------------------------------------------------------------------------
# Stub generator
# ---------------------------------------------------------------------------

CUSTOMERS_FILE = DATA / "graph" / "customers.json"


def approved_reference_customers() -> list[str]:
    """Same predicate as retriever._dynamo_reference_customers, read from the
    file that seeds the table instead of from DynamoDB."""
    today = date.today().isoformat()
    out = []
    for c in json.loads(CUSTOMERS_FILE.read_text()):
        expires = c.get("approval_expires") or ""
        if c.get("public_reference") and expires and expires >= today:
            out.append(c["name"])
    return out


def all_customer_names() -> list[str]:
    return [c["name"] for c in json.loads(CUSTOMERS_FILE.read_text())]


_DEFLECTION = (
    "Pricing and commercial terms are owned by the Deal Desk and are never generated "
    "automatically. This question has been routed to the commercial team."
)


def compose_answer(
    question: str,
    passages: list[RetrievedPassage],
    forced: bool,
    roster: list[str],
) -> GeneratedAnswer:
    """Extractive stand-in for the Bedrock generator.

    Restating the question is deliberate; see the module docstring.
    """
    if forced:
        text = f"{_DEFLECTION}\n\nRestating the question: {question}"
        return GeneratedAnswer(
            answer_text=text,
            citations=[],
            invoked_customers=[n for n in roster if n.lower() in text.lower()],
            model_id="stub-extractive",
            prompt_hash="stub",
            response_hash="stub",
        )

    primary = [p for p in passages if p.source_system in PRIMARY_SYSTEMS]
    grounding = (primary or passages)[:2]

    body = "\n\n".join(_first_sentences(p.excerpt, 3) for p in grounding)
    text = f"Restating the question: {question}\n\n{body}".strip()

    # Cite what grounded the answer, plus the strongest non-primary passage:
    # a generator that borrows phrasing from a prior questionnaire cites it,
    # and that second source_system is what lifts the coverage signal.
    cited = [p.document_id for p in grounding]
    for p in passages:
        if p.source_system not in PRIMARY_SYSTEMS and p.document_id not in cited:
            cited.append(p.document_id)
            break

    return GeneratedAnswer(
        answer_text=text,
        citations=cited,
        invoked_customers=[n for n in roster if n.lower() in text.lower()],
        model_id="stub-extractive",
        prompt_hash="stub",
        response_hash="stub",
    )


def _first_sentences(text: str, n: int) -> str:
    flat = " ".join(text.split())
    pieces = re.split(r"(?<=[.!?])\s+", flat)
    return " ".join(pieces[:n])


# ---------------------------------------------------------------------------
# One question through the pipeline
# ---------------------------------------------------------------------------


@dataclass
class Outcome:
    question_id: str
    text: str
    topics: list[str]
    topic_class: str
    tier: Tier
    scorer_tier: Tier
    composite: float
    breakdown: dict[str, float]
    triggers: list[str]
    citations: list[str]
    passages: list[RetrievedPassage]
    invoked_customers: list[str]
    expired_citations: list[str]


def run_question(qid: str, text: str, retriever: StubRetriever, roster: list[str]) -> Outcome:
    topics = _extract_topics(text)
    plan = get_dispatch_plan(topics)
    forced = plan.force_tier is not None

    passages: list[RetrievedPassage] = []
    if not forced:
        for source in list(plan.primary) + list(plan.secondary) + list(plan.tertiary):
            # deal_desk is deferred and seismic/gong are behind the mock-sources
            # API, which is not reachable offline. retriever/handler.py logs and
            # returns [] for both, so the offline result matches the deployed one
            # whenever that API is degraded - which is what questions 008-010
            # are there to prove.
            passages.extend(retriever.retrieve(text, source))

    prior = retriever.prior_matches(text) if plan.topic_class == "unclassified" else []

    generated = compose_answer(text, passages, forced, roster)
    composite, breakdown, scorer_tier = composite_score(
        passages=passages,
        prior_matches=prior,
        generated=generated,
        guardrail_flags=[],
        topic_class=plan.topic_class,
    )

    result = hard_rules.apply_rules(
        answer_text=generated.answer_text,
        invoked_customers=generated.invoked_customers,
        approved_reference_customers=approved_reference_customers(),
        current_tier=scorer_tier,
    )
    triggers = list(result.triggers)
    final = result.final_tier

    # hard_rules/handler.py applies the dispatch force_tier after the rules.
    if forced and plan.force_tier is Tier.RED:
        final = Tier.RED
        reason = plan.reason or "dispatch_force_tier"
        if reason not in triggers:
            triggers.append(reason)

    expired = [
        p.document_id
        for p in passages
        if p.provenance is not None and p.provenance.is_expired
    ]

    return Outcome(
        question_id=qid,
        text=text,
        topics=topics,
        topic_class=plan.topic_class,
        tier=final,
        scorer_tier=scorer_tier,
        composite=round(composite, 4),
        breakdown={k: round(v, 4) for k, v in breakdown.model_dump().items()},
        triggers=triggers,
        citations=generated.citations,
        passages=passages,
        invoked_customers=generated.invoked_customers,
        expired_citations=sorted(set(expired)),
    )


# ---------------------------------------------------------------------------
# Checks
# ---------------------------------------------------------------------------

EXPECTED_COUNTS = {"compliance": 6, "product-docs": 10, "prior-rfps": 6, "sme-approved": 2}


def check_corpus() -> tuple[bool, list[str]]:
    lines: list[str] = []
    ok = True
    total = 0
    for prefix, expected in EXPECTED_COUNTS.items():
        directory = CORPUS_REAL / prefix
        docs = sorted(p for p in directory.glob("*") if not p.name.endswith(".metadata.json"))
        missing = [d.name for d in docs if not d.with_name(d.name + ".metadata.json").exists()]
        orphans = [
            s.name for s in directory.glob("*.metadata.json")
            if not s.with_name(s.name[: -len(".metadata.json")]).exists()
        ]
        status = "OK "
        if len(docs) != expected or missing or orphans:
            ok = False
            status = "FAIL"
        lines.append(
            f"  {status} {prefix + '/':<16} {len(docs):>2} documents "
            f"(expected {expected}), {len(docs) - len(missing)} with sidecars"
        )
        if missing:
            lines.append(f"       documents without a .metadata.json: {missing}")
        if orphans:
            lines.append(f"       sidecars without a document: {orphans}")
        total += len(docs)
    lines.append(f"  total documents: {total} (expected {sum(EXPECTED_COUNTS.values())})")
    return ok and total == sum(EXPECTED_COUNTS.values()), lines


def check_parses(manifest: dict[str, Any]) -> tuple[bool, list[str]]:
    """The real workbook parser must find exactly the questions the manifest
    describes, in order. A fixture the parser mis-reads is worse than no
    fixture, because the failure shows up as missing rows in the demo."""
    from excel_parser.parser import parse_workbook  # noqa: PLC0415

    lines: list[str] = []
    ok = True
    _, parsed = parse_workbook(FIXTURE)
    expected_texts = [q["text"] for q in manifest["questions"]]
    got_texts = [q.text for q in parsed]
    if got_texts == expected_texts:
        lines.append(
            f"  OK   excel_parser found all {len(parsed)} questions in order, "
            f"across {len({q.section for q in parsed})} sections"
        )
    else:
        ok = False
        lines.append(
            f"  FAIL excel_parser found {len(parsed)} questions, manifest lists "
            f"{len(expected_texts)}"
        )
        for missing in set(expected_texts) - set(got_texts):
            lines.append(f"       not parsed: {missing[:70]}")
    return ok, lines


def check_tiers(
    outcomes: dict[str, Outcome], manifest: dict[str, Any]
) -> tuple[bool, list[str], list[str]]:
    lines: list[str] = []
    diverged: list[str] = []
    for q in manifest["questions"]:
        out = outcomes[q["id"]]
        if out.tier.value != q["expected_tier"]:
            diverged.append(
                f"  {q['id']}  expected {q['expected_tier']:<5} got {out.tier.value:<5} "
                f"[class={out.topic_class} composite={out.composite} "
                f"triggers={out.triggers or '-'}]"
            )
    if diverged:
        lines.append(f"  {len(diverged)} of {len(manifest['questions'])} questions diverge:")
        lines.extend(diverged)
    else:
        lines.append(f"  all {len(manifest['questions'])} questions reproduce the expected tier")

    missing_triggers: list[str] = []
    for q in manifest["questions"]:
        out = outcomes[q["id"]]
        absent = [t for t in q["expected_triggers"] if t not in out.triggers]
        if absent:
            missing_triggers.append(f"  {q['id']}  missing triggers {absent}, got {out.triggers}")
    if missing_triggers:
        lines.append(f"  {len(missing_triggers)} questions missing an expected trigger:")
        lines.extend(missing_triggers)
    else:
        lines.append("  every expected per-question trigger fired")

    return not diverged and not missing_triggers, lines, diverged


REQUIRED_TRIGGERS = [
    "pricing_detected",
    "compliance_claim",
    "sla_claim",
    "unapproved_reference:",  # prefix match — the customer name is appended
    "forward_looking_statement",
    "competitor_disparagement",
]


def check_triggers(outcomes: dict[str, Outcome]) -> tuple[bool, list[str]]:
    fired: dict[str, list[str]] = {}
    for out in outcomes.values():
        for t in out.triggers:
            fired.setdefault(t, []).append(out.question_id)

    lines: list[str] = []
    ok = True
    for required in REQUIRED_TRIGGERS:
        if required.endswith(":"):
            hits = {t: q for t, q in fired.items() if t.startswith(required)}
            if hits:
                for name, qids in sorted(hits.items()):
                    lines.append(f"  OK   {name:<42} {', '.join(qids)}")
            else:
                ok = False
                lines.append(f"  FAIL {required + '*':<42} never fired")
        elif required in fired:
            lines.append(f"  OK   {required:<42} {', '.join(fired[required])}")
        else:
            ok = False
            lines.append(f"  FAIL {required:<42} never fired")

    extra = sorted(set(fired) - set(REQUIRED_TRIGGERS) - {
        t for t in fired if t.startswith("unapproved_reference:")
    })
    if extra:
        lines.append(f"  also fired (not required): {', '.join(extra)}")
    return ok, lines


def check_spread(outcomes: dict[str, Outcome], floors: dict[str, int]) -> tuple[bool, list[str]]:
    counts = Counter(o.tier.value for o in outcomes.values())
    lines: list[str] = []
    ok = True
    for tier, floor in floors.items():
        got = counts.get(tier, 0)
        status = "OK " if got >= floor else "FAIL"
        if got < floor:
            ok = False
        lines.append(f"  {status} {tier.upper():<6} {got:>2} (floor {floor})")
    lines.append(f"  total {sum(counts.values())}")
    return ok, lines


def check_freshness(
    outcomes: dict[str, Outcome], manifest: dict[str, Any]
) -> tuple[bool, list[str]]:
    """Assert what the shipped code actually guarantees about expired evidence.

    Not what the brief assumed. retriever/handler.py removed query-time
    freshness suppression, so nothing demotes an expired primary during
    retrieval and nothing stops a generator grounding on it. What does hold,
    and what is asserted here:

      a. the superseded report is reachable in the corpus;
      b. Provenance.is_expired is True and render_citation() suffixes the
         source line with (EXPIRED), so a reviewer sees it;
      c. score_f() floors the superseded report at 0.2 while the current one
         scores 1.0, so the freshness signal separates them;
      d. the current report is retrieved alongside it on both questions, so
         the correct evidence is always available to the answer;
      e. the question is amber, so a human reviews it before it ships.

    Retrieval rank is reported, not asserted. On CAIQ-015 the superseded report
    legitimately outranks the current one, because the question asks about it
    by date. That is the honest behaviour and the reason (b) and (e) matter.
    """
    from confidence_scorer.scorer import score_f  # noqa: PLC0415

    lines: list[str] = []
    ok = True
    for spec in manifest.get("expected_suppressed_documents", []):
        doc_id = spec["document_id"]
        current_id = spec["superseded_by"]
        seen_expired = False

        for qid in spec["on_questions"]:
            out = outcomes[qid]
            ranked = sorted(out.passages, key=lambda p: p.score, reverse=True)
            comp = [p for p in ranked if p.source_system == "compliance"]
            expired_hits = [p for p in comp if p.document_id == doc_id]
            current_hits = [p for p in comp if p.document_id == current_id]

            # (d) the current report must always be available.
            if not current_hits:
                ok = False
                lines.append(f"  FAIL {qid}: current report {current_id} was not retrieved")
                continue

            # (c) the freshness signal must separate the two.
            f_current = score_f([current_hits[0]])
            lines.append(
                f"  OK   {qid}: {current_id} retrieved at rank "
                f"{ranked.index(current_hits[0]) + 1}/{len(ranked)}, "
                f"score {current_hits[0].score}, freshness F={f_current}"
            )
            if f_current < 1.0:
                ok = False
                lines.append(f"  FAIL {qid}: current report should score F=1.0, got {f_current}")

            if expired_hits:
                seen_expired = True
                p = expired_hits[0]
                f_expired = score_f([p])
                prov = p.provenance
                # (b) must be flagged, or a reviewer has no signal at all.
                if prov is None or not prov.is_expired:
                    ok = False
                    lines.append(
                        f"  FAIL {qid}: {doc_id} retrieved but NOT flagged expired "
                        f"(sidecar expires_on problem)"
                    )
                else:
                    citation = render_citation(prov)
                    if not citation.endswith("(EXPIRED)"):
                        ok = False
                        lines.append(f"  FAIL {qid}: citation lacks the (EXPIRED) suffix")
                    else:
                        rank = ranked.index(p) + 1
                        note = " — outranks the current report on this question" if (
                            rank < ranked.index(current_hits[0]) + 1
                        ) else ""
                        lines.append(
                            f"  OK   {qid}: {doc_id} at rank {rank}/{len(ranked)}, "
                            f"freshness F={f_expired}{note}"
                        )
                        lines.append(f"       citation reads: {citation}")
                if f_expired > 0.2:
                    ok = False
                    lines.append(
                        f"  FAIL {qid}: superseded report should floor at F=0.2, got {f_expired}"
                    )

            # (e) the question must not be green.
            if out.tier is Tier.GREEN:
                ok = False
                lines.append(f"  FAIL {qid}: tier is green; expired evidence would ship unreviewed")

        # (a) reachable somewhere across the pair.
        if not seen_expired:
            ok = False
            lines.append(
                f"  FAIL {doc_id} was never retrieved on {spec['on_questions']} — "
                f"the freshness demo has nothing to show"
            )
    return ok, lines


# ---------------------------------------------------------------------------
# Second run — the step-7 flywheel
# ---------------------------------------------------------------------------

# What an SME would type into the review UI for CAIQ-038, which run one leaves
# red with no corpus coverage. Held here rather than written into
# data/corpus/sme-approved/ on purpose: persisting it would make the corpus
# three SME answers instead of two, and would mean run one no longer starts
# from the state the demo starts from.
SME_ANSWER_038 = (
    "The service does not install an agent on managed devices, so there is nothing to "
    "deploy through unified endpoint management. Where a tenant uses the mobile client, "
    "it is distributed through the managed application catalogue of the tenant's UEM "
    "platform and accepts configuration through the AppConfig managed-configuration "
    "channel that iOS and Android expose, so tenant identifier, federation endpoint and "
    "policy pinning can be pushed without user entry. Per-app VPN, managed-open-in "
    "restrictions and remote wipe are enforced by the UEM platform against the client "
    "like any other managed application. For unmanaged mobile devices, recipients use "
    "the browser-based portal, which stores no message content on the device."
)


def simulate_followup(retriever: StubRetriever, roster: list[str], manifest: dict[str, Any]) -> list[str]:
    """Re-run the follow-up workbook after ingesting the SME's answer.

    Mirrors review_api.handler._emit_sme_approved_to_kb(): the document id is
    derived from a hash of the question text, source_type is
    sme_approved_answer, and the sidecar carries approved_at, approved_by and a
    two-year expires_on.
    """
    import hashlib  # noqa: PLC0415

    by_id = {q["id"]: q for q in manifest["questions"]}
    before = {
        qid: run_question(qid, by_id[qid]["text"], retriever, roster)
        for qid in [f["id"] for f in manifest["followup"]["questions"]]
    }

    q038 = by_id["CAIQ-038"]["text"]
    doc_id = "lfb-" + hashlib.sha256(q038.encode("utf-8")).hexdigest()[:16]
    approved_at = date.today().isoformat()
    expires_on = date(date.today().year + 2, date.today().month, date.today().day).isoformat()
    sidecar = {
        "document_id": doc_id,
        "source_type": "sme_approved_answer",
        "source_id": "sme_approved_answers",
        "topic_ids": [],
        "approved_at": approved_at,
        "approved_by": "r.calloway@echoworx.example",
        "expires_on": expires_on,
        "question_text": q038[:500],
    }
    text = f"{q038}\n\n{SME_ANSWER_038}"
    retriever.chunks.append(
        Chunk(
            document_id=doc_id,
            source_type="sme_approved_answer",
            heading="Question and answer",
            text=text,
            sidecar=sidecar,
            uri=f"s3://reference-corpus/sme-approved/{doc_id}.md",
            tokens=Counter(_tokens(text)),
        )
    )
    retriever.idf = build_idf(retriever.chunks)

    after = {
        qid: run_question(qid, by_id[qid]["text"], retriever, roster)
        for qid in before
    }

    lines = [
        f"  SME answer ingested as sme-approved/{doc_id}.md "
        f"(approved {approved_at}, expires {expires_on})",
        "",
        f"  {'id':<10}{'run 1':<22}{'run 2':<22}cites the SME answer",
    ]
    for qid, b in before.items():
        a = after[qid]
        cited = doc_id in [m for m in [a.citations] for m in m] or any(
            p.document_id == doc_id for p in a.passages
        )
        # sme_approved answers reach the scorer through prior_matches, not
        # passages, so ask the retriever directly whether it matched.
        matched = [
            m.answer_id
            for m in retriever.prior_matches(by_id[qid]["text"])
        ]
        lines.append(
            f"  {qid:<10}"
            f"{b.tier.value + ' H=' + format(b.breakdown['h'], '.3f') + ' c=' + format(b.composite, '.3f'):<22}"
            f"{a.tier.value + ' H=' + format(a.breakdown['h'], '.3f') + ' c=' + format(a.composite, '.3f'):<22}"
            f"{'yes' if (doc_id in matched or cited) else 'no'}"
        )
    return lines


# ---------------------------------------------------------------------------
# Fixture reading
# ---------------------------------------------------------------------------


def read_fixture() -> list[tuple[str, str]]:
    """Read (id, text) straight from the workbook, so the harness verifies the
    file the pipeline will actually receive and not the generator's in-memory
    list."""
    from openpyxl import load_workbook  # noqa: PLC0415 — optional at import time

    wb = load_workbook(FIXTURE, data_only=True)
    ws = wb.active
    out: list[tuple[str, str]] = []
    for row in ws.iter_rows(min_row=7, values_only=True):
        qid, text = (row + (None, None))[:2]
        if isinstance(qid, str) and isinstance(text, str) and qid.startswith("CAIQ-"):
            out.append((qid, text))
    return out


# ---------------------------------------------------------------------------


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--verbose", action="store_true", help="print per-question detail")
    ap.add_argument("--json", type=Path, help="write the full per-question result to this path")
    ap.add_argument(
        "--followup",
        action="store_true",
        help="simulate the second run: ingest the SME's answer to CAIQ-038 and re-score",
    )
    args = ap.parse_args()

    if not FIXTURE.exists() or not MANIFEST.exists():
        raise SystemExit("Fixture or manifest missing. Run scripts/create_echoworx_fixture.py.")

    manifest = json.loads(MANIFEST.read_text())
    retriever = StubRetriever(build_index())
    roster = all_customer_names()

    questions = read_fixture()
    outcomes = {
        qid: run_question(qid, text, retriever, roster) for qid, text in questions
    }

    print("=" * 78)
    print("Echoworx CAIQ fixture verification")
    print(f"  fixture        {FIXTURE.relative_to(ROOT)}  ({len(questions)} questions)")
    print(f"  rules version  {hard_rules.RULES_VERSION}")
    print(f"  corpus chunks  {len(retriever.chunks)}")
    print(f"  approved refs  {approved_reference_customers()}")
    print("=" * 78)

    results: list[tuple[str, bool, list[str]]] = []

    ok, lines = check_corpus()
    results.append(("1. Corpus documents and sidecars", ok, lines))

    parse_ok, parse_lines = check_parses(manifest)
    ok, lines, _ = check_tiers(outcomes, manifest)
    results.append(("2. Expected tier per question", ok and parse_ok, parse_lines + lines))

    ok, lines = check_triggers(outcomes)
    results.append(("3. Hard-rule trigger coverage", ok, lines))

    ok, lines = check_spread(outcomes, manifest["tier_spread_floor"])
    results.append(("4. Tier spread", ok, lines))

    ok, lines = check_freshness(outcomes, manifest)
    results.append(("5. Expired-report handling", ok, lines))

    all_ok = True
    for title, ok, lines in results:
        print(f"\n{title}: {'PASS' if ok else 'FAIL'}")
        for line in lines:
            print(line)
        all_ok = all_ok and ok

    if args.verbose:
        print("\nPer-question detail")
        print(
            f"  {'id':<10}{'tier':<7}{'class':<17}{'comp':<7}"
            f"{'H':<7}{'R':<7}{'C':<7}{'F':<7}topics / triggers"
        )
        for qid, _ in questions:
            o = outcomes[qid]
            b = o.breakdown
            print(
                f"  {qid:<10}{o.tier.value:<7}{o.topic_class:<17}{o.composite:<7}"
                f"{b['h']:<7}{b['r']:<7}{b['c']:<7}{b['f']:<7}"
                f"{','.join(o.topics) or '-'}  |  {','.join(o.triggers) or '-'}"
            )

    if args.followup:
        print("\nSecond run (step-7 flywheel), simulated")
        for line in simulate_followup(retriever, roster, manifest):
            print(line)

    if args.json:
        args.json.write_text(json.dumps(
            {
                "generated_at": datetime.now(UTC).isoformat(),
                "rules_version": hard_rules.RULES_VERSION,
                "results": [
                    {
                        "id": o.question_id,
                        "topics": o.topics,
                        "topic_class": o.topic_class,
                        "scorer_tier": o.scorer_tier.value,
                        "final_tier": o.tier.value,
                        "composite": o.composite,
                        "breakdown": o.breakdown,
                        "triggers": o.triggers,
                        "citations": o.citations,
                        "invoked_customers": o.invoked_customers,
                        "expired_passages": o.expired_citations,
                    }
                    for o in outcomes.values()
                ],
            },
            indent=2,
        ))
        print(f"\nWrote {args.json}")

    print("\n" + "=" * 78)
    print("RESULT:", "PASS" if all_ok else "FAIL")
    print("=" * 78)
    return 0 if all_ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
